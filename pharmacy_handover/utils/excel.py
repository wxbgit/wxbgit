# -*- coding: utf-8 -*-
"""Excel 备份生成工具（openpyxl）。"""
import os
import io
import re
import tempfile

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage

from config import EXPORT_DIR
from utils.crypto import decode_data_url

TITLE_FILL = PatternFill('solid', fgColor='1F6FB2')
HEAD_FILL = PatternFill('solid', fgColor='DCE9F5')
WARN_FILL = PatternFill('solid', fgColor='FFE0E0')
THIN = Side(style='thin', color='B0B0B0')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _safe(text):
    return re.sub(r'[\\/:*?"<>|]', '_', str(text or ''))


def _write_kv(ws, row, k, v, key_width=18):
    ws.cell(row=row, column=1, value=k).font = Font(bold=True)
    ws.cell(row=row, column=1).fill = HEAD_FILL
    ws.cell(row=row, column=1).border = BORDER
    c = ws.cell(row=row, column=2, value=v)
    c.border = BORDER
    c.alignment = Alignment(wrap_text=True, vertical='top')
    return row + 1


def _section(ws, row, title, span=6):
    ws.cell(row=row, column=1, value=title).font = Font(bold=True, color='FFFFFF', size=12)
    ws.cell(row=row, column=1).fill = TITLE_FILL
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    return row + 1


def build_excel(record):
    """根据交接班记录生成标准格式 Excel，返回文件路径。"""
    os.makedirs(EXPORT_DIR, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = '交接班记录'
    ws.column_dimensions['A'].width = 20
    for col in 'BCDEF':
        ws.column_dimensions[col].width = 20

    row = 1
    ws.cell(row=row, column=1, value='药师交接班记录表').font = Font(bold=True, size=16)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    ws.row_dimensions[row].height = 26
    row += 2

    row = _section(ws, row, '一、基础信息')
    row = _write_kv(ws, row, '记录编号', record.get('id'))
    row = _write_kv(ws, row, '交接时间', record.get('handover_time'))
    row = _write_kv(ws, row, '交班药师', '%s(%s)' % (record.get('handover_name', ''), record.get('handover_id', '')))
    row = _write_kv(ws, row, '接班药师', '%s(%s)' % (record.get('receive_name', ''), record.get('receive_id', '')))
    row = _write_kv(ws, row, '所在药房', record.get('department'))
    row += 1

    row = _section(ws, row, '二、环境信息')
    row = _write_kv(ws, row, '药房温度(℃)', record.get('temperature'))
    row = _write_kv(ws, row, '药房湿度(%)', record.get('humidity'))
    fridges = record.get('fridge_data') or {}
    if isinstance(fridges, dict):
        fridges = fridges.get('list', [])
    if fridges:
        ws.cell(row=row, column=1, value='冰箱名称').font = Font(bold=True)
        ws.cell(row=row, column=2, value='温度(℃)').font = Font(bold=True)
        ws.cell(row=row, column=3, value='湿度(%)').font = Font(bold=True)
        ws.cell(row=row, column=4, value='是否正常').font = Font(bold=True)
        for c in range(1, 5):
            ws.cell(row=row, column=c).fill = HEAD_FILL
            ws.cell(row=row, column=c).border = BORDER
        row += 1
        for f in fridges:
            ws.cell(row=row, column=1, value=f.get('name', '')).border = BORDER
            ws.cell(row=row, column=2, value=f.get('temperature', '')).border = BORDER
            ws.cell(row=row, column=3, value=f.get('humidity', '')).border = BORDER
            ok = f.get('normal', '')
            cell = ws.cell(row=row, column=4, value=ok)
            cell.border = BORDER
            if str(ok).strip() not in ('正常', '是', 'True', 'true', '1'):
                cell.fill = WARN_FILL
            row += 1
    row += 1

    row = _section(ws, row, '三、麻精药品')
    nar = record.get('narcotic_info') or {}
    for k, label in (('variety_count', '在库品种数'), ('stock_qty', '在库数量'),
                     ('empty_ampoule', '空安瓿回收数'), ('safe_status', '保险柜状态'),
                     ('remark', '备注')):
        row = _write_kv(ws, row, label, nar.get(k, ''))
    row += 1

    row = _section(ws, row, '四、普通药品')
    nor = record.get('normaldruginfo') or {}
    row = _write_kv(ws, row, '在库品种总数', nor.get('variety_total', ''))
    row = _write_kv(ws, row, '在库总数量', nor.get('stock_total', ''))
    row += 1

    row = _section(ws, row, '五、高警示药品')
    drugs = record.get('highalertdrugs') or {}
    if isinstance(drugs, dict):
        drugs = drugs.get('list', [])
    if drugs:
        for idx, name in enumerate(['药品名称', '规格', '在库数量', '单位', '存放位置'], start=1):
            ws.cell(row=row, column=idx, value=name).font = Font(bold=True)
            ws.cell(row=row, column=idx).fill = HEAD_FILL
            ws.cell(row=row, column=idx).border = BORDER
        row += 1
        for d in drugs:
            for idx, key in enumerate(('name', 'spec', 'qty', 'unit', 'location'), start=1):
                ws.cell(row=row, column=idx, value=d.get(key, '')).border = BORDER
            row += 1
    row += 1

    row = _section(ws, row, '六、发药机状态')
    mach = record.get('machine_status') or {}
    row = _write_kv(ws, row, '快速发药机余量(%)', mach.get('fast_machine', ''))
    row = _write_kv(ws, row, '智能药架余量(%)', mach.get('smart_shelf', ''))
    row = _write_kv(ws, row, '其他说明', mach.get('remark', ''))
    row += 1

    row = _section(ws, row, '七、待办事项')
    pend = record.get('pending_items') or {}
    for k, label in (('unreviewed_rx', '未审核处方数'), ('undispensed_rx', '未调配处方数'),
                     ('return_orders', '待退药单数'), ('pending_advice', '待处理医嘱数')):
        row = _write_kv(ws, row, label, pend.get(k, ''))
    row += 1

    row = _section(ws, row, '八、异常事件')
    abn = record.get('abnormal_events') or {}
    row = _write_kv(ws, row, '异常类型', '、'.join(abn.get('events', []) or []) or '无')
    row = _write_kv(ws, row, '异常描述', abn.get('detail', ''))
    row += 1

    row = _section(ws, row, '九、交接留言')
    ws.cell(row=row, column=1, value=record.get('handover_note', '')).alignment = Alignment(wrap_text=True)
    ws.merge_cells(start_row=row, start_column=1, end_row=row + 2, end_column=6)
    row += 4

    # 库存快照（近效期标红）
    items = record.get('items') or []
    if items:
        row = _section(ws, row, '十、库存快照（近效期≤90天标红）')
        headers = ['药品名称', '规格', '批号', '有效期', '在库数量', '单位', '类型', '存放位置']
        for idx, h in enumerate(headers, start=1):
            ws.cell(row=row, column=idx, value=h).font = Font(bold=True)
            ws.cell(row=row, column=idx).fill = HEAD_FILL
            ws.cell(row=row, column=idx).border = BORDER
        row += 1
        for it in items:
            values = [it.get('drug_name', ''), it.get('spec', ''), it.get('batch_no', ''),
                      it.get('expiry_date', ''), it.get('stock_qty', ''), it.get('unit', ''),
                      it.get('drug_type', ''), it.get('location', '')]
            for idx, v in enumerate(values, start=1):
                cell = ws.cell(row=row, column=idx, value=v)
                cell.border = BORDER
                if it.get('near_expiry'):
                    cell.fill = WARN_FILL
            row += 1
        row += 1

    # 签名图片
    row = _section(ws, row, '十一、双方签名')
    for label, key in (('交班药师签名', 'handover_sign'), ('接班药师签名', 'receive_sign')):
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        data = decode_data_url(record.get(key, ''))
        if data:
            tmp = os.path.join(tempfile.gettempdir(), 'sign_%s_%s.png' % (record.get('id'), key))
            with open(tmp, 'wb') as f:
                f.write(data)
            try:
                img = XLImage(tmp)
                img.width, img.height = 160, 60
                ws.add_image(img, 'B%d' % row)
                ws.row_dimensions[row].height = 50
            except Exception:
                ws.cell(row=row, column=2, value='(签名图片)')
        else:
            ws.cell(row=row, column=2, value='(无签名)')
        row += 2

    row = _section(ws, row, '十二、防篡改信息')
    row = _write_kv(ws, row, 'PDF SHA256', record.get('pdf_hash') or '(未生成PDF)')

    date_str = str(record.get('handover_time') or '')[:10].replace('-', '')
    fname = _safe('%s_%s_No%s.xlsx' % (date_str or 'nodate', record.get('department') or '药房', record.get('id')))
    path = os.path.join(EXPORT_DIR, fname)
    wb.save(path)
    return path


def build_history_excel(records, filename):
    """导出历史记录汇总 Excel。"""
    os.makedirs(EXPORT_DIR, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = '交接班历史'
    headers = ['编号', '交接时间', '交班药师', '接班药师', '药房', '温度', '湿度', '异常事件', '留言']
    for idx, h in enumerate(headers, start=1):
        ws.cell(row=1, column=idx, value=h).font = Font(bold=True)
        ws.cell(row=1, column=idx).fill = HEAD_FILL
        ws.cell(row=1, column=idx).border = BORDER
        ws.column_dimensions[get_column_letter(idx)].width = 16
    row = 2
    for r in records:
        abn = r.get('abnormal_events') or {}
        values = [r.get('id'), r.get('handover_time'), r.get('handover_name'), r.get('receive_name'),
                  r.get('department'), r.get('temperature'), r.get('humidity'),
                  '、'.join(abn.get('events', []) or []), r.get('handover_note', '')]
        for idx, v in enumerate(values, start=1):
            ws.cell(row=row, column=idx, value=v).border = BORDER
        row += 1
    path = os.path.join(EXPORT_DIR, _safe(filename))
    wb.save(path)
    return path
